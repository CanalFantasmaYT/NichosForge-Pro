"""Exportar datos a Excel (.xlsx) con hojas, encabezados con color y columnas ajustadas.
Si openpyxl no está instalado, se entrega CSV (separado por ; y con BOM para que Excel lea los acentos)."""
import csv
import io

from . import motor

COLS_CANAL = [("canal", "Canal"), ("url", "Link"), ("rol", "Rol"), ("idioma", "Idioma"), ("pais", "País"),
              ("categoria", "Categoría"), ("score", "Puntaje"), ("suscriptores", "Suscriptores"),
              ("visitas_totales", "Visitas totales"), ("visitas_mes", "Visitas/mes"), ("metodo_mes", "Cómo se midió"), ("visitas_7d", "Visitas 7 días (medidas)"),
              ("visitas_30d", "Visitas 30 días (medidas)"), ("crecimiento", "Crecimiento semanal %"), ("videos", "Videos"),
              ("edad_dias", "Edad (días)"), ("creado", "Creado"), ("promedio", "Visitas promedio"), ("mediana", "Mediana"),
              ("techo", "Techo"), ("techo_titulo", "Video techo"), ("piso", "Piso"), ("likes_prom", "Me gusta prom."),
              ("comentarios_prom", "Comentarios prom."), ("interaccion", "Interacción %"), ("vistas_por_sub", "Visitas/sub"),
              ("pct_largos", "% videos largos"), ("pct_shorts", "% Shorts"), ("videos_semana", "Videos/semana"),
              ("tipo", "Tipo"), ("auto_tipo", "Tipo de canal"), ("auto_score", "% automatizado"), ("rostro_v", "Rostro (verificado)"), ("monet_v", "¿Monetizado?"),
              ("audiencia", "Audiencia"), ("salud", "Estado"), ("cpm", "CPM est. $"), ("rpm", "RPM est. $"), ("ingreso_mes", "Ganancia/mes est. $"),
              ("ingreso_mes_bajo", "Ganancia/mes mín. $"), ("ingreso_mes_alto", "Ganancia/mes máx. $"),
              ("ingreso_anio", "Ganancia/año est. $"), ("es_referente", "Referente"), ("riesgos", "Riesgos"),
              ("tags_top", "Etiquetas principales")]
COLS_VIDEO = [("titulo", "Título"), ("url", "Link"), ("publicado", "Publicado"), ("vistas", "Visitas"), ("likes", "Me gusta"),
              ("comentarios", "Comentarios"), ("interaccion", "Interacción %"), ("vistas_dia", "Visitas/día"),
              ("outlier", "x su mediana"), ("duracion_seg", "Duración (s)"), ("tags", "Etiquetas")]
COLS_VIRAL = [("titulo", "Título"), ("canal", "Canal"), ("url", "Link"), ("publicado", "Publicado"), ("vistas", "Visitas"),
              ("likes", "Me gusta"), ("comentarios", "Comentarios"), ("vistas_dia", "Visitas/día"),
              ("outlier", "x su canal"), ("duracion_seg", "Duración (s)"), ("frase", "Encontrado con"), ("tags", "Etiquetas")]


def _filas_canal(filas):
    nombres = {"con_rostro": "Con rostro", "sin_rostro": "Sin rostro", "dudoso": "Dudoso"}
    return [{**f, "rostro_v": nombres.get((f.get("rostro") or {}).get("veredicto"), "Sin verificar"),
             "monet_v": (f.get("monetizado") or {}).get("titulo", "")} for f in filas]


def _val(v):
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    if isinstance(v, bool):
        return "Sí" if v else "No"
    return v


def _libro(hojas: list) -> tuple:
    """hojas = [(nombre, columnas[(clave, titulo)], filas[dict])] -> (bytes, extension, mime)"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        nombre, cols, filas = hojas[0]
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow([t for _, t in cols])
        for f in filas:
            w.writerow([_val(f.get(k)) for k, _ in cols])
        return ("\ufeff" + buf.getvalue()).encode("utf-8"), "csv", "text/csv; charset=utf-8"
    wb = Workbook()
    wb.remove(wb.active)
    relleno = PatternFill("solid", fgColor="4F7A63")
    for nombre, cols, filas in hojas:
        ws = wb.create_sheet(nombre[:31])
        ws.append([t for _, t in cols])
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), relleno, Alignment(vertical="center", wrap_text=True)
        for f in filas:
            ws.append([_val(f.get(k)) for k, _ in cols])
        ws.freeze_panes = "A2"
        for i, (k, t) in enumerate(cols, 1):
            largo = max([len(str(t))] + [min(60, len(str(_val(f.get(k)) or ""))) for f in filas[:200]])
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(55, largo + 2))
            if k in ("url", "techo_url"):
                for r in range(2, ws.max_row + 1):
                    cel = ws.cell(r, i)
                    if cel.value:
                        cel.hyperlink, cel.font = cel.value, Font(color="2F6FB0", underline="single")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _resumen_canal(a: dict) -> list:
    filas = [{"m": t, "v": _val(a.get(k))} for k, t in COLS_CANAL if k != "tags_top"]
    filas += [{"m": "Método visitas/mes", "v": a.get("metodo_mes")}, {"m": "Base del CPM", "v": a.get("base")}]
    return filas


def canal(cid: str):
    a = motor.canal(cid)
    if not a:
        raise ValueError("Canal no analizado.")
    vids = sorted(a.get("videos_lista", []), key=lambda v: v["vistas"], reverse=True)
    hojas = [("Resumen", [("m", "Métrica"), ("v", "Valor")], _resumen_canal(a)),
             ("Videos", COLS_VIDEO, vids),
             ("Etiquetas", [("tag", "Etiqueta"), ("usos", "Usos"), ("vistas_prom", "Visitas promedio")], a.get("etiquetas", [])),
             ("Por día", [("dia", "Día"), ("videos", "Videos"), ("vistas_prom", "Visitas promedio")], a.get("semana", [])),
             ("Por duración", [("tramo", "Duración"), ("videos", "Videos"), ("vistas_prom", "Visitas promedio")], a.get("duraciones", []))]
    if a.get("top_historico"):
        hojas.insert(2, ("Top histórico", COLS_VIDEO, a["top_historico"]))
    return _libro(hojas), a["canal"]


def busqueda(bid: str):
    b = motor.busqueda(bid)
    canales = [motor.ligero_de(cid) for cid in b["canales"]]
    canales = [c for c in canales if c]
    return _libro([("Canales", COLS_CANAL, _filas_canal(canales)), ("Videos virales", COLS_VIRAL, b.get("virales", []))]), b["titulo"]


COLS_INV = [("nicho", "Nicho"), ("titulo", "Título"), ("canal", "Canal"), ("url", "Link"), ("vistas", "Visitas"), ("vph", "Visitas/hora"),
            ("outlier", "x su canal"), ("likes", "Me gusta"), ("comentarios", "Comentarios"), ("publicado", "Publicado"),
            ("duracion_seg", "Duración (s)"), ("auto_tipo", "Tipo de canal"), ("auto_score", "% automatizado"), ("suscriptores", "Subs canal"),
            ("idioma", "Idioma")]


def investigacion(rid: str):
    r = motor.investigacion(rid)
    if not r:
        raise ValueError("Investigación no encontrada.")
    cols_c = [("nicho", "Nicho")] + COLS_CANAL
    return _libro([("Videos", COLS_INV, r["videos"]), ("Canales", cols_c, r["canales"]),
                   ("Nichos", [("n", "Nicho"), ("c", "Videos")], [{"n": n, "c": c} for n, c in r["nichos"]])]), r["titulo"]


def videos_guardados():
    filas = sorted(motor.videos_guardados().values(), key=lambda v: (v.get("categoria") or "", -(v.get("vistas") or 0)))
    cols = [("categoria", "Categoría"), ("notas", "Notas")] + COLS_INV[1:]
    return _libro([("Videos guardados", cols, filas)]), "videos guardados"


def guardados():
    filas = motor.lista_guardados()
    cols = [("etapa", "Etapa"), ("nicho", "Nicho"), ("notas", "Notas")] + COLS_CANAL
    return _libro([("Guardados", cols, _filas_canal(filas))]), "guardados"
